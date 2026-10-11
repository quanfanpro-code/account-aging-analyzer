# -*- coding: utf-8 -*-
import os
from datetime import date
import openpyxl
import pytest
from src.core.fifo_engine import FIFOEngine
from src.core.report import ReportGenerator


def test_生成Excel_包含七个Sheet(tmp_path):
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 365), (365, 730)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=1000.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 500.0, "credit": 0.0}],
    )
    output = tmp_path / "审计底稿.xlsx"
    gen = ReportGenerator()
    gen.generate(
        output_path=output,
        aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1),
                "buckets": [(0, 365), (365, 730)],
                "direction": "debit"},
    )
    assert output.exists()
    wb = openpyxl.load_workbook(output)
    assert "借方账龄汇总表" in wb.sheetnames
    assert "贷方账龄汇总表" in wb.sheetnames
    assert "成对对抵后账龄" in wb.sheetnames
    assert "逐笔追溯明细" in wb.sheetnames
    assert "期初余额验证" in wb.sheetnames
    assert "期末余额验证" in wb.sheetnames
    assert "参数说明" in wb.sheetnames
    assert len(wb.sheetnames) == 7


def test_账龄汇总表_含合计行(tmp_path):
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 365)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=0.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 1000.0, "credit": 0.0}],
    )
    output = tmp_path / "out.xlsx"
    ReportGenerator().generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1), "buckets": [(0, 365)], "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["借方账龄汇总表"]
    # 最后一行或前两行应有"合计"
    all_text = " ".join(
        str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
    )
    assert "合计" in all_text


def test_期初余额验证_含期初分档列(tmp_path):
    """当期初分档存在时，期初余额验证表应展示各分档"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 365), (365, 730)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=1000.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 500.0, "credit": 0.0}],
        opening_buckets=[400.0, 600.0],
        opening_date=date(2024, 1, 1),
    )
    output = tmp_path / "opening_test.xlsx"
    gen = ReportGenerator()
    gen.generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1),
                "buckets": [(0, 365), (365, 730)],
                "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["期初余额验证"]
    all_text = " ".join(
        str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
    )
    assert "期初分档" in all_text


def test_自定义分档标签_出现在Excel表头(tmp_path):
    """传入 bucket_labels 时，Excel 表头应使用自定义标签而非 '0-90天'"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 90), (90, 365), (365, 9999)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=1000.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 500.0, "credit": 0.0}],
    )
    output = tmp_path / "label_test.xlsx"
    gen = ReportGenerator()
    gen.generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1),
                "buckets": [(0, 90), (90, 365), (365, 9999)],
                "direction": "debit",
                "bucket_labels": ["3个月内", "4-12月", "1年以上"]},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["借方账龄汇总表"]
    # 表头应包含自定义标签
    header_row = [str(cell.value) for cell in sheet[1]]
    header_text = " ".join(header_row)
    assert "3个月内" in header_text
    assert "4-12月" in header_text
    assert "1年以上" in header_text
    # 不应包含原始的天数格式
    assert "0-90天" not in header_text


def test_无自定义标签时回退到天数格式(tmp_path):
    """不传 bucket_labels 时，表头仍用 '0-90天' 格式"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 90), (90, 365)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=0.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 1000.0, "credit": 0.0}],
    )
    output = tmp_path / "nolabel.xlsx"
    gen = ReportGenerator()
    gen.generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1), "buckets": [(0, 90), (90, 365)], "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["借方账龄汇总表"]
    header_text = " ".join(str(cell.value) for cell in sheet[1])
    assert "0-90天" in header_text
    assert "90-365天" in header_text


def test_参数说明_包含期初日期与模式说明(tmp_path):
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 90), (90, 365)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=1000.0,
        transactions=[],
        opening_date=date(2024, 1, 1),
    )
    output = tmp_path / "config.xlsx"
    ReportGenerator().generate(
        output_path=output,
        aging_results=[result],
        config={
            "cutoff_date": date(2025, 6, 1),
            "buckets": [(0, 90), (90, 365)],
            "direction": "debit",
            "opening_date": date(2024, 1, 1),
            "aging_mode": "铁板模式",
            "bucket_boundary_rule": "边界点归前档",
        },
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["参数说明"]
    values = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "期初余额日期" in values
    assert "铁板模式" in values
    assert "边界点归前档" in values


def test_账龄汇总表_合计列等于分档合计而非原始期末(tmp_path):
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 365), (365, 730)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=1000.0,
        transactions=[{"date": date(2024, 6, 1), "debit": -1500.0, "credit": 0.0}],
        opening_date=date(2024, 1, 1),
    )
    output = tmp_path / "negative.xlsx"
    ReportGenerator().generate(
        output_path=output,
        aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1), "buckets": [(0, 365), (365, 730)], "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["借方账龄汇总表"]
    header_text = " ".join(str(cell.value) for cell in sheet[1] if cell.value)
    assert "反向余额（待核对）" in header_text
    # B2 起写：列号偏移 +1（A=1 留空，B=2 起写）
    bucket_sum = (sheet.cell(row=2, column=5).value or 0) + (sheet.cell(row=2, column=6).value or 0)
    negative_value = sheet.cell(row=2, column=7).value or 0
    total = sheet.cell(row=2, column=8).value
    assert bucket_sum == pytest.approx(0.0)
    assert negative_value == pytest.approx(-500.0)
    assert total == pytest.approx(bucket_sum + negative_value)


def test_期末余额验证_含独立勾稽列(tmp_path):
    """期末余额验证表应有"分档合计+负余额"独立勾稽列，与期末余额对比"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 365), (365, 730)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=1000.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 500.0, "credit": 0.0}],
    )
    output = tmp_path / "closing_check.xlsx"
    ReportGenerator().generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1), "buckets": [(0, 365), (365, 730)],
                "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["期末余额验证"]
    header = [str(cell.value) for cell in sheet[1] if cell.value]
    assert "分档合计+负余额" in header
    assert "勾稽差异" in header
    # 数据行：分档合计+负余额 = total_check = closing_balance，勾稽差异 = 0
    idx = header.index("分档合计+负余额")
    check_val = sheet.cell(row=2, column=idx + 2).value  # B 列起写，偏移 +1
    assert check_val == pytest.approx(result.total_check)
    diff_val = sheet.cell(row=2, column=idx + 3).value
    assert diff_val == pytest.approx(0.0)


def test_合计行SUM公式写入正确sheet且字体为黑(tmp_path):
    """_patch_summary_formulas 应把 SUM 公式写进'借方账龄汇总表'，且公式单元格为黑字"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 365)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=0.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 1000.0, "credit": 0.0}],
    )
    output = tmp_path / "formula.xlsx"
    ReportGenerator().generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1), "buckets": [(0, 365)], "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["借方账龄汇总表"]
    # 1 个明细 → 合计行 = 第 3 行；分档 1 列 + 合计列 = E、F 两列写公式
    cell = sheet.cell(row=3, column=5)
    assert cell.value == "=SUM(E2:E2)"
    assert cell.font.color is not None and str(cell.font.color.rgb).endswith("000000")


def test_显式空白或重复账龄名称拒绝导出():
    with pytest.raises(ValueError, match="不能为空"):
        ReportGenerator._get_bucket_labels({
            "buckets": [(0, 90), (90, 180)], "bucket_labels": ["", ""],
        })
    with pytest.raises(ValueError, match="不能重复"):
        ReportGenerator._get_bucket_labels({
            "buckets": [(0, 90), (90, 180)], "bucket_labels": ["半年内", "半年内"],
        })


def test_未提供名称时按档位生成非空默认名称():
    assert ReportGenerator._get_bucket_labels({
        "buckets": [(0, 90), (730, 9999)],
    }) == ["0-90天", "730天以上"]


def test_无反向余额也保留反向余额列():
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1), buckets=[(0, 365)], direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=100.0, transactions=[],
    )
    summary = ReportGenerator()._build_summary(
        [result], {"buckets": [(0, 365)], "bucket_labels": ["1年以内"]},
    )
    assert list(summary.columns) == [
        "科目编码", "科目名称", "明细名称", "1年以内",
        "反向余额（待核对）", "合计",
    ]
    assert summary.iloc[0]["反向余额（待核对）"] == 0.0


def test_参数说明使用固定方向和反向余额新口径():
    config_table = ReportGenerator()._build_config({
        "buckets": [(0, 365)],
        "bucket_labels": ["1年以内"],
        "direction_policy": "八类往来按科目类别固定方向",
        "negative_balance_policy": "反向余额以负数单列展示，不改变科目所属汇总表",
    })
    text = " ".join(str(value) for value in config_table.to_numpy().ravel())
    assert "八类往来按科目类别固定方向" in text
    assert "反向余额以负数单列展示" in text
    assert "默认方向" not in text
