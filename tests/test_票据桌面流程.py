from datetime import date
import queue
import threading

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.core.current_accounts import filter_current_accounts
from src.core.column_mapper import ColumnMapper
from src.core.data_loader import DataLoader
from src.core.workflow import AgingWorkflow
from src.gui.step_result import StepResult
from src.gui.票据确认窗口 import collect_bill_decisions


def source(kind='商业承兑汇票'):
    return pd.DataFrame([
        [date(2023,1,1),'销售','1122','应收账款','客户甲',200,0,'',''],
        [date(2025,6,1),'收票','1122','应收账款','客户甲',0,200,'',''],
        [date(2025,6,1),'收票','1121','应收票据','客户甲',200,0,kind,'票甲'],
    ],columns=['日期','凭证号','科目编码','科目名称','明细名称','借方金额','贷方金额','票据类型','票据号'])


class Runner:
    _run_calculation=StepResult._run_calculation
    def __init__(self,frame):
        self._state={'journal_df':frame,'cutoff_date':date(2025,12,31),
            'buckets':[(0,365),(366,9999)],'include_receivable_bills':True}
        self._loader=DataLoader();self._msg_queue=queue.Queue();self.logs=[]
    def _log(self,message):self.logs.append(message)


def test_桌面后台实际计算票据并保存导出依据(tmp_path):
    runner=Runner(source());runner._run_calculation()
    result=runner._state['aging_results'][0]
    assert result.subject_name=='应收票据' and result.bucket_totals==[0,200]
    assert runner._state['calculation_trace']['bill_transfer_records'][0]['承接金额']==200
    output=tmp_path/'桌面票据账龄.xlsx'
    StepResult._do_export(runner,str(output))
    book=load_workbook(output,data_only=True)
    assert '票据连续计龄依据' in book.sheetnames and '票据口径确认' in book.sheetnames
    assert any('【票据连续计龄】' in str(cell.value) for row in book['账龄计算结果'] for cell in row)
    book.close()


@pytest.mark.parametrize('choice,expected',[('连续计龄',[0,200]),('不连续计龄',[200,0])])
def test_类型不明先请求确认再继续后台计算(choice,expected):
    runner=Runner(source('银行承兑'))
    worker=threading.Thread(target=runner._run_calculation,daemon=True);worker.start()
    while True:
        kind,payload=runner._msg_queue.get(timeout=5)
        if kind=='bill_confirmation':
            items,event,answer=payload
            assert 'aging_results' not in runner._state
            answer['decisions']=collect_bill_decisions(items,[choice],['按已核实合同与会计政策确定本项计龄口径'])
            event.set();break
        assert kind!='calc_error',payload
    worker.join(timeout=5)
    assert not worker.is_alive()
    assert runner._state['aging_results'][0].bucket_totals==expected
    assert runner._state['calculation_trace']['bill_policy_records'][0]['连续计龄'] is (choice=='连续计龄')


def test_取消票据确认不会留下本次成功结果():
    runner=Runner(source('银行承兑'))
    worker=threading.Thread(target=runner._run_calculation,daemon=True);worker.start()
    while True:
        kind,payload=runner._msg_queue.get(timeout=5)
        if kind=='bill_confirmation':
            _,event,answer=payload;answer['decisions']=None;event.set();break
        assert kind!='calc_error',payload
    worker.join(timeout=5)
    assert not worker.is_alive() and 'aging_results' not in runner._state


@pytest.mark.parametrize('choice,reason',[('未确认','依据'),('连续计龄','')])
def test_口径与依据缺一不可(choice,reason):
    with pytest.raises(ValueError):
        collect_bill_decisions([{'单位':'甲','金额':100}],[choice],[reason])


def test_导入可选识别票据而原八类默认不变():
    frame=source()
    assert filter_current_accounts(frame,data_kind='journal',source_label='测试').retained_rows==2
    report=filter_current_accounts(frame,data_kind='journal',source_label='测试',include_receivable_bills=True)
    assert report.retained_rows==3 and report.category_counts['应收票据']==1
    mapping=ColumnMapper().auto_map(['票据号码','汇票类型','货币名称'])
    assert mapping['票据号']=='票据号码' and mapping['票据类型']=='汇票类型' and mapping['币种']=='货币名称'


def test_只有科目名称也能纳入票据():
    frame=source().drop(columns=['科目编码'])
    opening=pd.DataFrame(columns=['科目名称','明细名称','借方余额','贷方余额'])
    result=AgingWorkflow(date(2025,12,31),[(0,365),(366,9999)]).run(opening,frame,include_receivable_bills=True)
    assert len(result)==1 and result[0].subject_name=='应收票据'


@pytest.mark.parametrize('code,name,expected',[('1121','应收账款','应收账款'),('1122','应收票据','应收票据')])
def test_名称编码冲突按名称且不得重复纳入(code,name,expected):
    frame=pd.DataFrame([[date(2025,1,1),code,name,'客户甲',100,0]],
        columns=['日期','科目编码','科目名称','明细名称','借方金额','贷方金额'])
    opening=pd.DataFrame(columns=['科目编码','科目名称','明细名称','借方余额','贷方余额'])
    result=AgingWorkflow(date(2025,12,31),[(0,365),(366,9999)]).run(opening,frame,include_receivable_bills=True)
    assert len(result)==1 and result[0].subject_name==expected and result[0].closing_balance==100


def test_明示跨日转票据须报边界而不是重置账龄():
    frame=source()
    frame.loc[1,'摘要']='应收转票据'
    frame.loc[2,'日期']=date(2025,6,2)
    opening=pd.DataFrame(columns=['科目编码','科目名称','明细名称','借方余额','贷方余额'])
    with pytest.raises(ValueError,match='跨日关联当前不支持'):
        AgingWorkflow(date(2025,12,31),[(0,365),(366,9999)]).run(opening,frame,include_receivable_bills=True)
