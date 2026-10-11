# -*- coding: utf-8 -*-
"""生成真实多工作表样例，验证八类往来全自动化主流程。"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import load_workbook

from src.core.current_accounts import filter_current_accounts, find_category_conflicts
from src.core.data_loader import DataLoader
from src.core.workflow import AgingWorkflow
from src.core.report import ReportGenerator


def main() -> int:
    categories = [
        ("1122", "应收账款", "debit"), ("1123", "预付账款", "debit"),
        ("1221", "其他应收款", "debit"), ("1486", "合同资产", "debit"),
        ("2202", "应付账款", "credit"), ("2203", "预收账款", "credit"),
        ("2241", "其他应付款", "credit"), ("2204", "合同负债", "credit"),
    ]
    with TemporaryDirectory() as temporary:
        root = Path(temporary)
        opening_path = root / "期初余额.xlsx"
        journal_path = root / "序时账.xlsx"
        output_path = root / "往来账龄底稿.xlsx"

        opening = pd.DataFrame({
            "科目编码": [item[0] for item in categories],
            "科目名称": [item[1] for item in categories],
            "明细名称": [None] + [f"往来单位{i}" for i in range(2, 9)],
            "借方余额": [100.0, 200.0, 300.0, 400.0, 0.0, 0.0, 0.0, 0.0],
            "贷方余额": [500.0, 0.0, 0.0, 0.0, 500.0, 600.0, 700.0, 800.0],
        })
        journal = pd.DataFrame({
            "日期": [date(2026, 1, 1)] * 10,
            "科目编码": [item[0] for item in categories] + ["6602", "1122"],
            "科目名称": [item[1] for item in categories] + ["管理费用", "应收账款坏账准备"],
            "明细名称": ["客户A"] + [f"往来单位{i}" for i in range(2, 9)] + ["部门", "客户A"],
            "借方金额": [0.0] * 10,
            "贷方金额": [0.0] * 10,
            "摘要": ["集成验证"] * 10,
        })
        with pd.ExcelWriter(opening_path) as writer:
            pd.DataFrame({"说明": ["封面"]}).to_excel(writer, sheet_name="说明", index=False)
            opening.to_excel(writer, sheet_name="期初", index=False)
        with pd.ExcelWriter(journal_path) as writer:
            pd.DataFrame({"说明": ["封面"]}).to_excel(writer, sheet_name="说明", index=False)
            journal.to_excel(writer, sheet_name="序时账", index=False)

        loader = DataLoader()
        assert loader.initial_sheet(loader.list_sheet_names(str(journal_path))) is None
        opening_raw = loader.load_sheet(str(opening_path), "期初")
        journal_raw = loader.load_sheet(str(journal_path), "序时账")
        print("SHEET_SELECTION_OK")

        opening_result = filter_current_accounts(
            opening_raw, data_kind="opening", source_label="期初余额/期初",
        )
        journal_result = filter_current_accounts(
            journal_raw, data_kind="journal", source_label="序时账/序时账",
        )
        assert opening_result.retained_rows == 8
        assert journal_result.retained_rows == 8
        assert journal_result.ignored_rows == 2
        assert not find_category_conflicts([opening_result, journal_result])
        print("CURRENT_ACCOUNT_FILTER_OK")

        for result in (opening_result, journal_result):
            assert result.original_rows == result.retained_rows + result.ignored_rows
            assert result.ignored_rows == sum(result.ignored_counts.values())
            assert result.retained_rows == sum(result.category_counts.values())
        print("FILTER_LOG_RECONCILIATION_OK")

        workflow = AgingWorkflow(
            cutoff_date=date(2026, 12, 31),
            buckets=[(0, 365)],
            opening_date=date(2025, 12, 31),
        )
        results = workflow.run(opening_result.dataframe, journal_result.dataframe)
        assert {result.direction for result in results} == {"debit", "credit"}
        assert next(
            result for result in results
            if result.subject_code == "2203" and result.detail_name == "（未填明细）"
        ).closing_balance == 400.0
        print("AUTO_RECLASSIFICATION_OK")

        config = {
            "cutoff_date": date(2026, 12, 31),
            "buckets": [(0, 365)],
            "bucket_labels": ["1年以内"],
            "opening_date": date(2025, 12, 31),
            "direction_policy": "反向余额自动重分类",
            "negative_balance_policy": "按转入科目追溯账龄",
            "source_results": workflow.source_results,
            "reclassification_records": workflow.reclassification_records,
            "offset_records": workflow.offset_records,
        }
        ReportGenerator().generate(output_path, results, config)
        workbook = load_workbook(output_path, data_only=False)
        assert "账龄计算结果" in workbook.sheetnames
        assert "成对对抵后账龄" not in workbook.sheetnames
        for sheet_name in ("账龄计算结果",):
            headers = [cell.value for cell in workbook[sheet_name][1] if cell.value]
            assert "1年以内" in headers
            assert "反向余额（待核对）" not in headers
            assert all(str(header).strip() for header in headers)
        print("BUCKET_HEADERS_OK")
        print("REVERSE_BALANCE_OK")

        assert all(abs(result.closing_balance - result.total_check) < 0.005 for result in results)
        print("RECONCILIATION_OK")
        print("INTEGRATION_OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
