# -*- coding: utf-8 -*-
"""生成深海蓝标准格式审计底稿"""
from __future__ import annotations
from collections import defaultdict
from pathlib import Path
from typing import Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from make_excel import make_excel
from src.core.current_accounts import BLANK_DETAIL, classify_current_account
from src.core.models import AgingResult

PathLike = Union[str, Path]

PAIRED_ACCOUNTS = (
    ("应收账款", "预收账款"),
    ("预付账款", "应付账款"),
    ("其他应收款", "其他应付款"),
    ("合同资产", "合同负债"),
)


class ReportGenerator:
    """生成账龄审计底稿 Excel（含截止日后的成对对抵展示）。"""

    def generate(
        self,
        output_path: PathLike,
        aging_results: list[AgingResult],
        config: dict,
    ) -> None:
        make_excel(self.build_frames(aging_results, config), output_path)

    def build_frames(self, aging_results: list[AgingResult], config: dict):
        """结果页只有最终账龄；其他页保存计算依据，不再输出并列账龄。"""
        source = config.get('source_results', aging_results)
        sheets = [
            ('账龄计算结果', self._build_final_summary(aging_results, config)),
            ('逐笔追溯明细', self._build_entries(aging_results, config)),
            ('原账勾稽', self._build_closing_check(source).drop(columns=['分档合计+负余额','勾稽差异'])),
        ]
        for title,key,columns in [
            ('处理确认记录','confirmation_records',['类型','原名称','用户决定','依据']),
            ('重分类记录','reclassification_records',['单位','原科目','原余额','转入科目','重分类金额','说明']),
            ('对抵记录','offset_records',['往来组合','单位','资产侧余额','负债侧余额','对抵金额','最终借方净额']),
        ]:
            records=config.get(key,[])
            sheets.append((title,pd.DataFrame(records,columns=columns)))
        for title,key in [('计算用期初','opening_source'),('计算用序时账','journal_source')]:
            if isinstance(config.get(key),pd.DataFrame):sheets.append((title,config[key]))
        sheets.append(('参数说明',self._build_config(config)))
        return sheets

    def _build_final_summary(self, results: list[AgingResult], config: dict):
        labels=self._get_bucket_labels(config)
        headers=['科目编码','科目名称','明细名称',*labels,'期末余额','列示类别']
        rows=[]
        for r in sorted(results,key=lambda x:(x.direction,x.subject_code,x.detail_name)):
            if r.closing_balance < -.005 or abs(r.negative_balance)>.005:
                raise ValueError('反向余额须先经统一计算流程自动重分类后再导出')
            if abs(r.closing_balance)<.005:continue
            if abs(sum(r.bucket_totals)-r.closing_balance)>.01:
                raise ValueError('最终账龄分档与期末余额不一致')
            rows.append([r.subject_code,r.subject_name,r.detail_name,*r.bucket_totals,r.closing_balance,'资产' if r.direction=='debit' else '负债'])
        # 资产与负债分别汇总，避免相加形成无业务含义的一个总数。
        for direction,label in [('debit','资产合计'),('credit','负债合计')]:
            selected=[r for r in results if r.direction==direction and abs(r.closing_balance)>=.005]
            if selected:rows.append([label,'','',*[sum(r.bucket_totals[i] for r in selected) for i in range(len(labels))],sum(r.closing_balance for r in selected),''])
        return pd.DataFrame(rows,columns=headers)

    # ─── 账龄汇总表 ───────────────────────────────────────────
    def _build_summary(self, results: list[AgingResult], config: dict) -> pd.DataFrame:
        labels = self._get_bucket_labels(config)

        headers = ["科目编码", "科目名称", "明细名称"]
        for label in labels:
            headers.append(label)
        headers.append("反向余额（待核对）")
        headers.append("合计")

        rows = []
        for r in results:
            row = [r.subject_code, r.subject_name, r.detail_name]
            for amt in r.bucket_totals:
                row.append(amt)
            row.append(r.negative_balance)
            row.append(r.total_check)
            rows.append(row)

        # 合计行：先写真实合计值（回退模式下这就是最终值）；
        # make_excel 模式下 _patch_summary_formulas 会再替换为 SUM 公式
        if results:
            sum_row = ["合计", "", ""]
            for i in range(len(results[0].bucket_totals)):
                sum_row.append(sum(r.bucket_totals[i] for r in results))
            sum_row.append(sum(r.negative_balance for r in results))
            sum_row.append(sum(r.total_check for r in results))
            rows.append(sum_row)
        else:
            # 无数据时加占位行，确保 sheet 被创建
            placeholder = ["无数据", "", ""]
            for label in labels:
                placeholder.append(0.0)
            placeholder.append(0.0)
            placeholder.append(0.0)
            rows.append(placeholder)

        return pd.DataFrame(rows, columns=headers)

    def _build_paired_offset(self, results: list[AgingResult], config: dict) -> pd.DataFrame:
        """按同一单位和成对科目展示净额；对抵额先冲减较老的剩余账龄。"""
        labels = self._get_bucket_labels(config)
        headers = [
            "往来组合", "明细名称", "资产侧科目", "负债侧科目", "资产侧期末余额",
            "负债侧期末余额", "对抵金额", "对抵后方向", "对抵后余额",
            *labels, "分档勾稽差异", "处理状态",
        ]
        category_side = {
            category: (pair_index, side)
            for pair_index, pair in enumerate(PAIRED_ACCOUNTS)
            for side, category in enumerate(pair)
        }
        grouped: dict[tuple[int, str], list[list[AgingResult]]] = defaultdict(lambda: [[], []])
        for result in results:
            match = classify_current_account(result.subject_code, result.subject_name)
            if match.category not in category_side:
                continue
            pair_index, side = category_side[match.category]
            grouped[(pair_index, result.detail_name)][side].append(result)

        rows = []
        for (pair_index, detail), (asset_results, liability_results) in sorted(grouped.items()):
            asset_name, liability_name = PAIRED_ACCOUNTS[pair_index]
            asset_amount = round(sum(r.closing_balance for r in asset_results), 2)
            liability_amount = round(sum(r.closing_balance for r in liability_results), 2)
            source_results = asset_results + liability_results
            can_offset = (
                detail != BLANK_DETAIL
                and bool(str(detail).strip())
                and all(r.closing_balance >= -0.005 for r in source_results)
                and all(len(r.bucket_totals) == len(labels) for r in source_results)
                and all(abs(sum(r.bucket_totals) - r.closing_balance) < 0.01 for r in source_results)
            )
            if not can_offset:
                status = "未填明细，未自动对抵" if detail == BLANK_DETAIL or not str(detail).strip() else "反向余额或分档待核对，未自动对抵"
                rows.append([
                    f"{asset_name}↔{liability_name}", detail, asset_name, liability_name,
                    asset_amount, liability_amount, 0.0, "待核对", None,
                    *([None] * len(labels)), None, status,
                ])
                continue

            offset = round(min(asset_amount, liability_amount), 2)
            net = round(asset_amount - liability_amount, 2)
            if abs(net) < 0.005:
                direction, balance, buckets = "结清", 0.0, [0.0] * len(labels)
            else:
                remaining_results = asset_results if net > 0 else liability_results
                direction = "借方" if net > 0 else "贷方"
                balance = abs(net)
                buckets = [round(sum(r.bucket_totals[i] for r in remaining_results), 2) for i in range(len(labels))]
                to_offset = offset
                for i in range(len(buckets) - 1, -1, -1):
                    consumed = min(buckets[i], to_offset)
                    buckets[i] = round(buckets[i] - consumed, 2)
                    to_offset = round(to_offset - consumed, 2)
            rows.append([
                f"{asset_name}↔{liability_name}", detail, asset_name, liability_name,
                asset_amount, liability_amount, offset, direction, balance,
                *buckets, round(sum(buckets) - balance, 2),
                "已对抵" if offset else ("无期末余额" if balance == 0 else "单侧余额"),
            ])
        return pd.DataFrame(rows, columns=headers)

    def _patch_summary_formulas(self, output_path: PathLike,
                                results: list[AgingResult], config: dict,
                                sheet_name: str) -> None:
        """在 make_excel 生成后，用 SUM 公式替换合计行的占位值"""
        if not results:
            return
        wb = load_workbook(output_path)
        ws = wb[sheet_name]

        labels = self._get_bucket_labels(config)
        sum_row = len(results) + 2  # make_excel 从第 1 行写表头，数据从第 2 行开始

        # 数字列从第 5 列（科目编码=B, 科目名称=C, 明细名称=D 之后）
        # make_excel 从 B 列(col=2) 开始写，所以：
        # B=科目编码, C=科目名称, D=明细名称, E=分档1, F=分档2, ..., 最后=合计
        first_num_col = 5  # E 列（B=科目编码, C=科目名称, D=明细名称 之后）
        last_num_col = 5 + len(labels) + 1  # 固定反向余额列之后为合计列

        for col in range(first_num_col, last_num_col + 1):
            col_letter = get_column_letter(col)
            cell = ws.cell(row=sum_row, column=col)
            cell.value = f"=SUM({col_letter}2:{col_letter}{sum_row - 1})"
            # 项目约定：蓝字=手动输入，黑字=公式；写公式后显式置黑
            cell.font = Font(name=cell.font.name, size=cell.font.size,
                             bold=cell.font.bold, color="000000")

        wb.save(output_path)

    # ─── 逐笔追溯明细 ───────────────────────────────────────────
    def _build_entries(self, results: list[AgingResult], config: dict) -> pd.DataFrame:
        headers = ["科目编码", "科目名称", "明细名称", "来源", "来源日期", "金额", "账龄天数", "分档下标"]
        rows = []
        for result in results:
            for entry in result.entries:
                rows.append([
                    result.subject_code, result.subject_name, result.detail_name,
                    entry.source, entry.source_date, entry.amount,
                    entry.days, entry.bucket_index,
                ])
        return pd.DataFrame(rows, columns=headers)

    # ─── 期初余额验证 ───────────────────────────────────────────
    def _build_opening_check(self, results: list[AgingResult]) -> pd.DataFrame:
        has_buckets = any(r.opening_bucket_totals for r in results)

        if has_buckets:
            bucket_count = max(len(r.opening_bucket_totals) for r in results)
            headers = ["科目", "明细", "期初余额"]
            for i in range(bucket_count):
                headers.append(f"期初分档{i+1}")
            headers.append("期初分档合计")
        else:
            headers = ["科目", "明细", "期初余额"]
            bucket_count = 0

        rows = []
        for r in results:
            row = [r.subject_name, r.detail_name, r.opening_balance]
            for i in range(bucket_count):
                row.append(r.opening_bucket_totals[i] if i < len(r.opening_bucket_totals) else 0)
            if has_buckets:
                row.append(sum(r.opening_bucket_totals))
            rows.append(row)

        return pd.DataFrame(rows, columns=headers)

    # ─── 期末余额验证 ───────────────────────────────────────────
    def _build_closing_check(self, results: list[AgingResult]) -> pd.DataFrame:
        # "分档合计+负余额"是独立勾稽列：与"验证期末"公式不同源，避免循环论证
        headers = ["科目", "明细", "方向", "期初余额", "借方合计", "贷方合计",
                   "验证期末", "引擎期末", "差异", "分档合计+负余额", "勾稽差异"]
        rows = []
        for r in results:
            if r.direction == "debit":
                calc = r.opening_balance + r.total_debit - r.total_credit
            else:
                calc = r.opening_balance + r.total_credit - r.total_debit
            rows.append([
                r.subject_name, r.detail_name,
                "借方" if r.direction == "debit" else "贷方",
                r.opening_balance, r.total_debit, r.total_credit,
                calc, r.closing_balance, r.closing_balance - calc,
                r.total_check, r.closing_balance - r.total_check,
            ])
        return pd.DataFrame(rows, columns=headers)

    # ─── 参数说明 ───────────────────────────────────────────
    def _build_config(self, config: dict) -> pd.DataFrame:
        labels = self._get_bucket_labels(config)
        rows = [
            ("计算截止日期", str(config.get("cutoff_date", ""))),
            ("账龄分档", ", ".join(labels)),
            ("方向口径", config.get("direction_policy", "反向余额自动重分类，按转入科目的方向计算账龄")),
            ("计算口径", config.get("fifo_policy", "先进先出法 (收付款按最老余额优先核销)")),
            ("期初余额日期", str(config.get("opening_date", ""))),
            ("期初模式", config.get("aging_mode", "铁板模式")),
            ("分档边界规则", config.get("bucket_boundary_rule", "边界点归前档")),
            ("反向余额说明", config.get(
                "negative_balance_policy",
                "反向余额自动重分类，并按转入科目的方向重新追溯计算账龄",
            )),
            ("最终结果口径", "名称确认与暂估归类后计算，反向余额自动重分类，再按成对科目对抵；最终仅列一套账龄，先扣较老余额"),
            ("未填明细提示", config.get(
                "blank_detail_policy",
                "归入（未填明细）的记录可能混合多个对手方，借贷互相核销后账龄结构可能失真，需人工评估",
            )),
        ]
        return pd.DataFrame(rows, columns=["参数", "值"])

    @staticmethod
    def _get_bucket_labels(config: dict) -> list[str]:
        """获取严格非空且唯一的账龄名称。"""
        buckets: list[tuple[int, int]] = config.get("buckets", [])
        if "bucket_labels" not in config or config.get("bucket_labels") is None:
            return [f"{lo}天以上" if hi == 9999 else f"{lo}-{hi}天" for lo, hi in buckets]
        labels = [str(label).strip() for label in config["bucket_labels"]]
        if len(labels) != len(buckets):
            raise ValueError("账龄名称数量必须与账龄档位数量一致")
        if any(not label for label in labels):
            raise ValueError("账龄名称不能为空")
        if len(set(labels)) != len(labels):
            raise ValueError("账龄名称不能重复")
        return labels
