# -*- coding: utf-8 -*-
"""报告分借/贷 sheet 测试"""
import os
from openpyxl import load_workbook
from src.core.report import ReportGenerator
from src.core.models import AgingResult


def test_报告分借贷两个sheet():
    results = [
        AgingResult(
            subject_code="1122", subject_name="应收账款",
            detail_name="客户A", direction="debit",
            opening_balance=1000.0, total_debit=500.0, total_credit=0.0,
            closing_balance=1500.0, bucket_totals=[1500.0],
        ),
        AgingResult(
            subject_code="2202", subject_name="应付账款",
            detail_name="供应商X", direction="credit",
            opening_balance=800.0, total_debit=200.0, total_credit=300.0,
            closing_balance=900.0, bucket_totals=[900.0],
        ),
    ]
    gen = ReportGenerator()
    output = "test_split_output.xlsx"
    gen.generate(output, results, {
        "cutoff_date": "2025-06-01",
        "buckets": [(0, 365)],
        "bucket_labels": ["1年以内"],
    })
    wb = load_workbook(output)
    assert "借方账龄汇总表" in wb.sheetnames
    assert "贷方账龄汇总表" in wb.sheetnames
    # 借方 sheet 只有 1122
    ws_d = wb["借方账龄汇总表"]
    # 找到科目编码列的值
    found_1122 = False
    found_2202 = False
    for row in ws_d.iter_rows(values_only=True):
        if "1122" in str(row):
            found_1122 = True
        if "2202" in str(row):
            found_2202 = True
    assert found_1122
    assert not found_2202
    # 贷方 sheet 只有 2202
    ws_c = wb["贷方账龄汇总表"]
    found_2202_c = False
    found_1122_c = False
    for row in ws_c.iter_rows(values_only=True):
        if "2202" in str(row):
            found_2202_c = True
        if "1122" in str(row):
            found_1122_c = True
    assert found_2202_c
    assert not found_1122_c
    # 清理
    os.remove(output)


def test_无数据方向仍生成工作表和提示行(tmp_path):
    result = AgingResult(
        subject_code="1122", subject_name="应收账款",
        detail_name="客户A", direction="debit",
        opening_balance=100.0, total_debit=0.0, total_credit=0.0,
        closing_balance=100.0, bucket_totals=[100.0],
    )
    output = tmp_path / "无贷方数据.xlsx"
    ReportGenerator().generate(output, [result], {
        "cutoff_date": "2025-06-01",
        "buckets": [(0, 365)],
        "bucket_labels": ["1年以内"],
    })
    workbook = load_workbook(output)
    assert "借方账龄汇总表" in workbook.sheetnames
    assert "贷方账龄汇总表" in workbook.sheetnames
    assert any(
        cell.value == "无数据"
        for row in workbook["贷方账龄汇总表"].iter_rows()
        for cell in row
    )
