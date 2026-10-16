from datetime import date

import openpyxl
import pandas as pd
import pytest

from src.core.models import AgingResult
from src.core.report import ReportGenerator
from src.core.workflow import AgingWorkflow, PAIRS, RULES


@pytest.mark.parametrize('original,target',[(a,b) for pair in PAIRS for a,b in (pair,pair[::-1])])
def test_四对往来双向重分类在主表末列注明(original,target):
    code,direction=RULES[target]
    result=AgingResult(code,target,'单位甲',direction,0,0,0,120,bucket_totals=[120])
    config={'buckets':[(0,365)],'reclassification_records':[
        {'单位':'单位甲','原科目':original,'原余额':-120,'转入科目':target,'重分类金额':120}]}
    frame=ReportGenerator()._build_final_summary([result],config)
    assert frame.columns[-1]=='备注'
    assert '【反向余额重分类】' in frame.iloc[0]['备注']
    assert original in frame.iloc[0]['备注'] and target in frame.iloc[0]['备注']
    assert '120.00' in frame.iloc[0]['备注']
    assert frame.iloc[-1]['备注']==''


def test_重分类来源缺显式记录时从原账结果识别且不跨往来组合():
    source=AgingResult('2202','应付账款','单位甲','credit',0,120,0,-120,negative_balance=-120)
    transferred=AgingResult('1123','预付账款','单位甲','debit',0,0,0,120,bucket_totals=[120])
    other=AgingResult('1221','其他应收款','单位甲','debit',0,0,0,30,bucket_totals=[30])
    frame=ReportGenerator()._build_final_summary([transferred,other],{'buckets':[(0,365)],'source_results':[source]})
    assert '应付账款' in frame[frame['科目名称']=='预付账款'].iloc[0]['备注']
    assert frame[frame['科目名称']=='其他应收款'].iloc[0]['备注']==''


def test_对抵后剩另一侧也保留重分类标注并醒目导出(tmp_path):
    opening=pd.DataFrame(columns=['科目编码','科目名称','明细名称','借方余额','贷方余额'])
    journal=pd.DataFrame([
        [date(2025,1,1),'1122','应收账款','单位甲',0,100,'预收'],
        [date(2025,2,1),'2203','预收账款','单位甲',150,0,'退付'],
    ],columns=['日期','科目编码','科目名称','明细名称','借方金额','贷方金额','摘要'])
    flow=AgingWorkflow(date(2025,12,31),[(0,365)])
    results=flow.run(opening,journal)
    target=tmp_path/'重分类备注.xlsx'
    ReportGenerator().generate(target,results,{'buckets':[(0,365)],'reclassification_records':flow.reclassification_records,'offset_records':flow.offset_records})
    book=openpyxl.load_workbook(target)
    sheet=book['账龄计算结果']
    assert sheet.cell(1,sheet.max_column).value=='备注'
    note=sheet.cell(2,sheet.max_column)
    assert '100.00' in note.value and '150.00' in note.value and '对抵' in note.value
    assert note.font.bold and note.font.color.rgb=='009C0006'
    assert note.alignment.wrap_text
    assert sheet['F2'].value==50
    book.close()
