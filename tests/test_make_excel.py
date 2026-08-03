# -*- coding: utf-8 -*-
"""make_excel 输出层测试"""
import pandas as pd
from openpyxl import load_workbook
from make_excel import make_excel


def test_空DataFrame也建Sheet并写占位行(tmp_path):
    """空 DataFrame 不应被静默跳过：建 sheet 并写"（无数据）"占位"""
    out = tmp_path / "empty.xlsx"
    make_excel(
        [("有数据", pd.DataFrame({"金额": [1.0]})),
         ("空表", pd.DataFrame({"金额": []}))],
        str(out),
    )
    wb = load_workbook(out)
    assert "有数据" in wb.sheetnames
    assert "空表" in wb.sheetnames
    texts = [str(c.value) for row in wb["空表"].iter_rows() for c in row if c.value]
    assert any("无数据" in t for t in texts)
