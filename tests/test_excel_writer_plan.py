# -*- coding: utf-8 -*-
"""postprocess 计划回调测试：回调收到表计划，改动落到输出文件。"""
import pandas as pd
from openpyxl import load_workbook
from make_excel import make_excel


def test_postprocess_收到表计划并可设页签色(tmp_path):
    plans = []

    def 记录(pl):
        plans.extend(pl)
        pl[0].标签色 = '0070C0'

    out = make_excel(pd.DataFrame({"金额": [1.0]}), str(tmp_path / "o.xlsx"),
                     postprocess=记录)
    assert len(plans) == 1            # 回调确实收到了表计划
    wb = load_workbook(out)
    assert wb["Sheet1"].sheet_properties.tabColor.rgb.endswith("0070C0")


def test_postprocess_区域覆盖落到单元格样式(tmp_path):
    def 高亮(pl):
        pl[0].apply_overrides({"B2": {"底色": "FFF2CC", "字色": "9C0006", "加粗": True}})

    out = make_excel(pd.DataFrame({"备注": ["【反向余额重分类】测试"]}), str(tmp_path / "o.xlsx"),
                     postprocess=高亮)
    wb = load_workbook(out)
    cell = wb["Sheet1"]["B2"]
    assert cell.fill.fgColor.rgb.endswith("FFF2CC")
    assert cell.font.bold is True
